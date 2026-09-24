"""Read the named source/oracle sheets. Output stays in the staging process."""
import datetime
import json
import sys

import openpyxl


def normalize(value):
    if isinstance(value, (datetime.date, datetime.datetime, datetime.time)):
        return value.isoformat()
    return value


def records(book, name, header_row=1):
    sheet = book[name]
    iterator = sheet.iter_rows(values_only=True)
    for _ in range(header_row - 1):
        next(iterator)
    keys = next(iterator)
    return [
        {key: normalize(value) for key, value in zip(keys, row) if key is not None}
        for row in iterator
        if any(value is not None for value in row)
    ]


source = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
master = openpyxl.load_workbook(sys.argv[2], read_only=True, data_only=True)
print(json.dumps({
    "source": {name: records(source, name) for name in (
        "Source_records", "Purchase_events", "Purchase_funding_components",
        "Supabase_economic_component_classifications",
    )},
    "master": {name: records(master, name, 4) for name in (
        "20_A2_EVENT_MAP", "21_A2_MONTH_STREAM", "22_A2_AMBIGUOUS",
        "23_A3_EVENT_LEDGER", "24_A3_MIXED_MATCH", "25_A3_DELTA_ORACLE",
        "26_A3_ANTI_DOUBLE", "27_A3_CONFLICTS",
    )},
}, ensure_ascii=False, separators=(",", ":")))
